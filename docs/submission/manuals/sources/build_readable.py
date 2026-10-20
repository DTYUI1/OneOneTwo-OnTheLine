#!/usr/bin/env python3
"""Readable manuals. Sources: Markdown. Build: python-docx + LibreOffice."""
from pathlib import Path
import argparse, re, subprocess, tempfile, json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader

DOCS = [
 ('01_technical_description','Техническое описание','Что это за тренажёр и как устроено обучение'),
 ('02_administrator_guide','Руководство администратора','Повседневная работа ответственного сотрудника класса'),
 ('03_teacher_guide','Руководство преподавателя','От первого занятия до выдачи результатов'),
 ('04_trainee_guide','Руководство обучаемого','Карточка, учебный телефон и разбор действий')]
BLUE='16324F'

def font(run, size=None, bold=False, color=None):
 run.font.name='Arial'; run.bold=bold
 run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:cs'),'Arial')
 if size: run.font.size=Pt(size)
 if color: run.font.color.rgb=RGBColor.from_string(color)

def inline(p,text):
 for part in re.split(r'(\*\*.*?\*\*|«[^»]+»)',text):
  if part: font(p.add_run(part.strip('*')),bold=part.startswith(('**','«')))

def fld(p,instruction):
 r=p.add_run(); font(r,11)
 a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin')
 b=OxmlElement('w:instrText'); b.set(qn('xml:space'),'preserve'); b.text=' '+instruction+' '
 c=OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'),'end')
 r._r.extend([a,b,c])

def mark(p,idx):
 a=OxmlElement('w:bookmarkStart'); a.set(qn('w:id'),str(idx)); a.set(qn('w:name'),f'ch{idx}')
 b=OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'),str(idx)); p._p.insert(0,a); p._p.append(b)

def link(p,label,idx):
 a=OxmlElement('w:hyperlink'); a.set(qn('w:anchor'),f'ch{idx}')
 r=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); co=OxmlElement('w:color'); co.set(qn('w:val'),BLUE); pr.append(co); r.append(pr)
 t=OxmlElement('w:t'); t.text=label; r.append(t); a.append(r); p._p.append(a)


def build(src,dst,title,subtitle):
 lines=src.read_text().splitlines(); doc=Document(); s=doc.sections[0]
 s.page_width=Cm(21); s.page_height=Cm(29.7)
 s.top_margin=Cm(1.9); s.bottom_margin=Cm(1.9); s.left_margin=Cm(2); s.right_margin=Cm(2)
 s.header_distance=Cm(.8); s.footer_distance=Cm(.8); s.different_first_page_header_footer=True
 normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(14)
 normal.paragraph_format.line_spacing=1.16; normal.paragraph_format.space_after=Pt(8)
 normal.paragraph_format.widow_control=True
 for name,size in [('Title',30),('Heading 1',21),('Heading 2',17),('Heading 3',15)]:
  st=doc.styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(BLUE)
  st.paragraph_format.keep_with_next=True; st.paragraph_format.space_before=Pt(13); st.paragraph_format.space_after=Pt(8)
 doc.styles['List Bullet'].font.name='Arial'; doc.styles['List Bullet'].font.size=Pt(14)
 hp=s.header.paragraphs[0]; font(hp.add_run('112  |  '+title),11,color=BLUE)
 fp=s.footer.paragraphs[0]; font(fp.add_run('Редакция 20.10.2026     •     Страница '),11); fld(fp,'PAGE'); font(fp.add_run(' из '),11); fld(fp,'NUMPAGES')
 doc.core_properties.title=title+' — тренажёр 112'; doc.core_properties.author='Проект 112'
 p=doc.add_paragraph(); p.paragraph_format.space_before=Cm(1.7); font(p.add_run('112 / УЧЕБНЫЙ ТРЕНАЖЁР'),15,True,BLUE)
 p=doc.add_paragraph(title,'Title'); p.paragraph_format.space_before=Cm(1.5)
 p=doc.add_paragraph(subtitle); font(p.runs[0],18,color=BLUE)
 p=doc.add_paragraph('Практическое пособие\n20 октября 2026 года'); p.paragraph_format.space_before=Cm(1.2)
 p=doc.add_paragraph('В тренажёре нет реальных вызовов.\nВсе звонки и происшествия — учебные.'); p.paragraph_format.space_before=Cm(1)
 doc.add_page_break(); p=doc.add_paragraph('Как пользоваться пособием','Heading 1')
 doc.add_paragraph('Названия кнопок и полей выделены полужирным и заключены в кавычки: «Войти». Номер шага означает одно действие. После действий прочитайте, что должно получиться.')
 doc.add_paragraph('Не обязательно читать всё подряд. Найдите нужную задачу в содержании. В электронной версии нажмите на название раздела, чтобы перейти к нему.')
 doc.add_paragraph('Если экран отличается от описания, не подбирайте случайные кнопки. Обратитесь к ответственному за обучение. Настройки и доступные возможности могут различаться.')
 for line in lines:
  if line.startswith('## '): break
  if line.strip() and not line.startswith('# '):
   p=doc.add_paragraph(); inline(p,line.strip())
 p=doc.add_paragraph('Содержание','Heading 1'); p.paragraph_format.page_break_before=True
 headings=[x[3:].strip() for x in lines if x.startswith('## ')]
 for idx,h in enumerate(headings,1):
  p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6)
  p.paragraph_format.tab_stops.add_tab_stop(Cm(17),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
  link(p,h,idx); p.add_run('\t'); fld(p,f'PAGEREF ch{idx} \\h')
 idx=0; first=True
 for line in lines:
  t=line.strip()
  if not t or t.startswith('# '): continue
  if first and not t.startswith('## '): continue
  if t.startswith('## '):
   idx+=1; p=doc.add_paragraph(t[3:],'Heading 1'); p.paragraph_format.page_break_before=(idx==1); mark(p,idx); first=False
  elif t.startswith('### '): doc.add_paragraph(t[4:],'Heading 2')
  elif t.startswith('#### '): doc.add_paragraph(t[5:],'Heading 3')
  elif t.startswith(('- ','* ')):
   p=doc.add_paragraph(style='List Bullet'); inline(p,t[2:].replace('[ ] ','□ ').replace('[x] ','☑ '))
  else:
   p=doc.add_paragraph(); inline(p,t.lstrip('> '))
   if re.match(r'^\d+[.)] ',t):
    p.paragraph_format.left_indent=Cm(.65); p.paragraph_format.first_line_indent=Cm(-.65); p.paragraph_format.keep_together=True
   if t in ('**Сделайте**','**Проверьте**','**Если не получилось**'):
    p.paragraph_format.keep_with_next=True
   if t.startswith(('Важно:','Проверьте:','Если не получилось:','Результат:')):
    pr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EEF3F7'); pr.append(sh)
    p.paragraph_format.space_before=Pt(5)
   if first: p.paragraph_format.space_after=Pt(12)
 # Keep a short closing chapter together instead of leaving a lone final note.
 last_heading=max(i for i,p in enumerate(doc.paragraphs) if p.style.name=='Heading 1')
 closing=doc.paragraphs[last_heading:]
 if sum(len(p.text) for p in closing)<2400:
  for p in closing[:-1]: p.paragraph_format.keep_with_next=True
  for p in closing: p.paragraph_format.keep_together=True
 settings=doc.settings.element; up=OxmlElement('w:updateFields'); up.set(qn('w:val'),'true'); settings.append(up)
 dst.parent.mkdir(parents=True,exist_ok=True); doc.save(dst)


def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--root',type=Path,default=Path(__file__).parent.parent); ap.add_argument('--fixture',action='store_true'); args=ap.parse_args()
 root=args.root; out=root/('fixture' if args.fixture else 'docx'); out.mkdir(parents=True,exist_ok=True)
 pdfdir=root/('fixture' if args.fixture else 'pdf'); pdfdir.mkdir(parents=True,exist_ok=True)
 specs=DOCS
 if args.fixture:
  src=out/'fixture.md'; src.write_text('# Проверка\n\nЁжик и съёмка: это тест вёрстки, не пособие.\n\n## 1. Проверка кириллицы и шага\n\n### Начало\n\n1. Нажмите «Проверить микрофон».\n\n2. Скажите: «Проверка связи».\n\nПроверьте: видны ё, Ё и цифры 123.\n\n## 2. Длинный текст\n\n'+('\n\n'.join(['Это проверочный абзац. '+ 'Читаемый текст и правильное расположение строк. '*12]*8)))
  build(src,out/'fixture.docx','Тест вёрстки — не пособие','Кириллица, содержание, нумерация, размер текста')
  paths=[out/'fixture.docx']
 else:
  paths=[]
  for stem,title,sub in specs:
   dst=out/(stem+'.docx'); build(root/'sources'/(stem+'.md'),dst,title,sub); paths.append(dst)
 with tempfile.TemporaryDirectory(prefix='lo-',dir=root) as tmp:
  p=subprocess.run(['/Applications/LibreOffice.app/Contents/MacOS/soffice','-env:UserInstallation='+Path(tmp).as_uri(),'--headless','--convert-to','pdf:writer_pdf_Export','--outdir',str(pdfdir),*[str(x) for x in paths]],capture_output=True,text=True,timeout=240)
  print(p.stdout); print(p.stderr)
  if p.returncode: raise RuntimeError('LibreOffice failed')
 records=[]
 qa=root/'verification'; qa.mkdir(exist_ok=True)
 for dx in paths:
  pdf=pdfdir/(dx.stem+'.pdf'); r=PdfReader(pdf); texts=[p.extract_text() or '' for p in r.pages]
  assert all(t.strip() for t in texts),'Empty PDF page'
  assert '\ufffd' not in ''.join(texts),'Replacement glyph'
  assert len(r.pages)>=3,'Short PDF'
  if not args.fixture:
   for heading in [x[3:].strip() for x in (root/'sources'/(dx.stem+'.md')).read_text().splitlines() if x.startswith('## ')]:
    normalized=lambda t: re.sub(r'\s+','',t)
    assert normalized(heading) in normalized('\n'.join(texts)),heading
  links=sum(len(p.get('/Annots',[])) for p in r.pages)
  assert links>0,'No contents links'
  fonts=set(); embedded=True
  for page in r.pages:
   for f in (page['/Resources'].get('/Font',{}) or {}).values():
    f=f.get_object(); fonts.add(str(f.get('/BaseFont','')))
    desc=f.get('/FontDescriptor')
    if desc:
     dd=desc.get_object(); embedded &= any(k in dd for k in ('/FontFile','/FontFile2','/FontFile3'))
    elif f.get('/DescendantFonts'):
     for d in f['/DescendantFonts']:
      dd=d.get_object()['/FontDescriptor'].get_object(); embedded &= any(k in dd for k in ('/FontFile','/FontFile2','/FontFile3'))
    else: embedded=False
  assert embedded,'Unembedded font'
  (qa/(dx.stem+'.txt')).write_text('\n\n'.join(texts))
  records.append(dict(file=str(pdf.relative_to(root)),pages=len(r.pages),bytes=pdf.stat().st_size,sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),fonts=sorted(fonts),all_fonts_embedded=embedded,links=links))
 manifest=dict(date='2026-10-20',font_size_pt=14,documents=records,verification_boundary='Content review and document rendering; no new application runtime acceptance')
 (qa/('fixture-manifest.json' if args.fixture else 'manifest.json')).write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
 print(json.dumps(manifest,ensure_ascii=False,indent=2))
if __name__=='__main__': main()
